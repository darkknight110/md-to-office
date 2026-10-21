#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""公文模板引擎（L2）的单元测试。

用一个**自己造的**空白公文模板做夹具：版心/页边距/字号/行距都按 GB/T 9704-2012 摆好，
但内容全是占位符——不含任何真实红头或单位信息。
"""
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import docx  # noqa: E402
from docx.enum.text import WD_LINE_SPACING  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402

from md_to_office import gongwen  # noqa: E402


def build_template(path, body_placeholder="{{正文}}"):
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(37), Mm(35)
    sec.left_margin, sec.right_margin = Mm(28), Mm(26)

    normal = d.styles["Normal"]
    normal.font.size = Pt(16)
    normal.font.name = "仿宋"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(28)
    pf.first_line_indent = Pt(32)  # 2 字符 @16pt

    h1 = d.styles["Heading 1"]
    h1.font.size = Pt(22)
    h1.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "方正小标宋简体")

    d.add_paragraph("{{标题}}", style="Heading 1")
    d.add_paragraph("{{主送机关}}")
    d.add_paragraph(body_placeholder)
    d.add_paragraph("{{成文日期}}")
    d.save(path)
    return path


class TestTemplateLifecycle(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = os.path.join(self.tmp.name, "store")
        self.tpl = build_template(os.path.join(self.tmp.name, "tpl.docx"))

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, *extra):
        return gongwen.main(["--store", self.store, "template", "add", self.tpl, "--name", "demo", *extra])

    def test_add_提取出版心与占位符(self):
        self.add()
        meta = gongwen.load_meta("demo", self.store)
        self.assertEqual(sorted(meta["placeholders"]), sorted(["标题", "主送机关", "正文", "成文日期"]))
        self.assertEqual(meta["params"]["margins_mm"], {"top": 37.0, "bottom": 35.0, "left": 28.0, "right": 26.0})
        self.assertEqual(meta["params"]["text_width_mm"], 156.0)
        self.assertEqual(meta["params"]["text_height_mm"], 225.0)
        self.assertEqual(meta["styles"]["Normal"]["size_pt"], 16.0)
        self.assertEqual(meta["styles"]["Normal"]["cjk_font"], "仿宋")
        self.assertEqual(meta["styles"]["Normal"]["line_spacing"], 28.0, "固定行距应按 pt 记，不是 EMU 原值")

    def test_add_拒绝重复登记(self):
        self.add()
        with self.assertRaises(SystemExit):
            self.add()

    def test_show_打印关键参数(self):
        self.add()
        # show 只打印，不抛异常
        self.assertEqual(gongwen.main(["--store", self.store, "template", "show", "demo"]), 0)

    def test_render_填充并保留样式(self):
        self.add()
        data = os.path.join(self.tmp.name, "d.json")
        import json
        json.dump({"标题": "关于测试的通知", "主送机关": "各处室：", "正文": "第一段。", "成文日期": "2026年10月21日"},
                  open(data, "w", encoding="utf-8"), ensure_ascii=False)
        out = os.path.join(self.tmp.name, "out.docx")
        rc = gongwen.main(["--store", self.store, "render", "--template", "demo", "--data", data, "--out", out])
        self.assertEqual(rc, 0)
        d = docx.Document(out)
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("关于测试的通知", text)
        self.assertNotIn("{{", text)
        sec = d.sections[0]
        self.assertAlmostEqual(sec.top_margin.mm, 37, places=1)
        self.assertEqual(d.styles["Normal"].font.size.pt, 16.0)

    def test_render_缺字段必须拒绝(self):
        self.add()
        out = os.path.join(self.tmp.name, "out.docx")
        with self.assertRaises(SystemExit):
            gongwen.main(["--store", self.store, "render", "--template", "demo", "--field", "标题=仅标题", "--out", out])

    def test_render_多段正文按段插入(self):
        self.add()
        data = os.path.join(self.tmp.name, "d.json")
        import json
        json.dump({"标题": "T", "主送机关": "A：", "正文": "第一段\n第二段\n第三段", "成文日期": "2026年10月21日"},
                  open(data, "w", encoding="utf-8"), ensure_ascii=False)
        out = os.path.join(self.tmp.name, "out.docx")
        gongwen.main(["--store", self.store, "render", "--template", "demo", "--data", data, "--out", out])
        d = docx.Document(out)
        texts = [p.text for p in d.paragraphs]
        for t in ("第一段", "第二段", "第三段"):
            self.assertIn(t, texts)
        # 插入段落必须继承 Normal 的 16pt / 28pt 行距
        idx = texts.index("第二段")
        self.assertEqual(d.paragraphs[idx].style.name, "Normal")

    def test_模板被改过必须提示(self):
        self.add()
        # 真的改内容（追加一段），确保 hash 变化；只重存一次相同内容 hash 不会变
        d = docx.Document(os.path.join(self.store, "demo", "template.docx"))
        d.add_paragraph("模板被改过")
        d.save(os.path.join(self.store, "demo", "template.docx"))
        out = os.path.join(self.tmp.name, "out.docx")
        with self.assertRaises(SystemExit):
            gongwen.main(["--store", self.store, "render", "--template", "demo", "--field", "标题=x",
                          "--field", "主送机关=y", "--field", "正文=z", "--field", "成文日期=w", "--out", out])

    def test_check_能发现参数被改(self):
        self.add()
        out = os.path.join(self.tmp.name, "out.docx")
        gongwen.main(["--store", self.store, "render", "--template", "demo",
                      "--field", "标题=x", "--field", "主送机关=y", "--field", "正文=z",
                      "--field", "成文日期=w", "--out", out])
        self.assertEqual(gongwen.main(["--store", self.store, "check", out, "--template", "demo"]), 0)

        # 人为把页边距改掉，check 必须失败
        d = docx.Document(out)
        d.sections[0].top_margin = Mm(20)
        d.save(out)
        self.assertEqual(gongwen.main(["--store", self.store, "check", out, "--template", "demo"]), 1)

    def test_check_能发现行距被改(self):
        self.add()
        out = os.path.join(self.tmp.name, "out2.docx")
        gongwen.main(["--store", self.store, "render", "--template", "demo",
                      "--field", "标题=x", "--field", "主送机关=y", "--field", "正文=z",
                      "--field", "成文日期=w", "--out", out])
        d = docx.Document(out)
        d.styles["Normal"].paragraph_format.line_spacing = Pt(20)
        d.save(out)
        self.assertEqual(gongwen.main(["--store", self.store, "check", out, "--template", "demo"]), 1)


class TestNoPlaceholder(unittest.TestCase):
    def test_无占位符模板会被提示(self):
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "plain.docx")
            d = docx.Document()
            d.add_paragraph("这是一份没有任何占位符的普通文档")
            d.save(p)
            store = os.path.join(td, "store")
            gongwen.main(["--store", store, "template", "add", p, "--name", "plain"])
            meta = gongwen.load_meta("plain", store)
            self.assertEqual(meta["placeholders"], [])


if __name__ == "__main__":
    unittest.main()


class TestLeftoverPlaceholder(unittest.TestCase):
    def test_允许缺省时提示残留(self):
        with tempfile.TemporaryDirectory() as td:
            tpl = build_template(os.path.join(td, "tpl.docx"))
            store = os.path.join(td, "store")
            gongwen.main(["--store", store, "template", "add", tpl, "--name", "t"])
            out = os.path.join(td, "out.docx")
            rc = gongwen.main(["--store", store, "render", "--template", "t",
                               "--field", "标题=x", "--allow-missing", "--out", out])
            self.assertEqual(rc, 0)
            text = "\n".join(p.text for p in docx.Document(out).paragraphs)
            self.assertIn("{{主送机关}}", text, "未替换的占位符应原样保留，便于用户发现")


class TestBodyParagraphs(unittest.TestCase):
    def test_正文空行不产生空段落(self):
        with tempfile.TemporaryDirectory() as td:
            tpl = build_template(os.path.join(td, "tpl.docx"))
            store = os.path.join(td, "store")
            gongwen.main(["--store", store, "template", "add", tpl, "--name", "t"])
            out = os.path.join(td, "out.docx")
            gongwen.main(["--store", store, "render", "--template", "t",
                          "--field", "标题=x", "--field", "主送机关=y",
                          "--field", "正文=第一段\n\n第二段", "--field", "成文日期=z", "--out", out])
            texts = [p.text for p in docx.Document(out).paragraphs]
            self.assertIn("第一段", texts)
            self.assertIn("第二段", texts)
            self.assertNotIn("", texts[1:-1], "正文里的空行不应变成空段落")
