#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""公文模板引擎（L2：模板固化）。

定位要说清楚：**它不是"生成公文"，是"按你给的模板填内容"**。
模板里的版心、字号、行距、页码、红头、版记、印章位都是你的，工具只负责：
  1. 读出来（inspect）并记住（add）；
  2. 把内容填进占位符，**保留模板的样式**（render）；
  3. 生成后量一遍，证明没跑偏（check）。

三条硬规则：
  · 印章/签名**绝不自动生成**，只保留位置；
  · 模板被换过（hash 变了）必须提示，不静默沿用旧参数；
  · 字体本机缺失时明确报告（Word 会静默替换成宋体 —— 那是最危险的一种"看起来对"）。

用法：
    python3 scripts/gongwen.py template add 模板.docx --name 某单位
    python3 scripts/gongwen.py template show 某单位
    python3 scripts/gongwen.py render --template 某单位 --data data.json --body 正文.md --out 成文.docx
    python3 scripts/gongwen.py check 成文.docx --template 某单位
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import re
import shutil
import sys
from datetime import datetime, timezone

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401
from docx.shared import Mm, Pt

PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def default_store() -> str:
    return os.environ.get("GONGWEN_HOME") or os.path.join(os.path.expanduser("~"), ".md-to-office", "templates")


def sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


# ─────────────────────────── 检查 ───────────────────────────

# 常见中文字体 → 各平台上的字体文件。命中就敢说"有/没有"，不命中才退化为"不确定"。
FONT_FILES = {
    "宋体": ["/System/Library/Fonts/Supplemental/Songti.ttc", "C:/Windows/Fonts/simsun.ttc",
             "/usr/share/fonts/truetype/arphic/uming.ttc"],
    "仿宋": ["/System/Library/Fonts/Supplemental/STFangsong.ttf", "C:/Windows/Fonts/simfang.ttf",
             "/usr/share/fonts/truetype/arphic/ukai.ttc"],
    "楷体": ["/System/Library/Fonts/Supplemental/Kaiti.ttc", "C:/Windows/Fonts/simkai.ttf"],
    "黑体": ["/System/Library/Fonts/STHeiti Light.ttc", "C:/Windows/Fonts/simhei.ttf"],
    "微软雅黑": ["C:/Windows/Fonts/msyh.ttc"],
    "方正小标宋简体": ["C:/Windows/Fonts/FZXBSJW.ttf"],
    "苹方": ["/System/Library/Fonts/PingFang.ttc"],
}


def font_available(name: str) -> bool | None:
    """本机是否有这个字体。

    返回 False 是**有依据的"没有"**（常见字体，按系统文件判断）；
    返回 None 才是"无法判断"——不假装知道，也不把不知道说成没有。
    """
    if not name:
        return None
    if name in FONT_FILES:
        return any(os.path.exists(p) for p in FONT_FILES[name])
    try:
        import subprocess
        out = subprocess.run(["fc-list", ":lang=zh", "family"], capture_output=True, text=True, timeout=5)
        if out.returncode == 0:
            return name.lower() in out.stdout.lower()
    except (FileNotFoundError, subprocess.SubprocessError):
        pass
    return None


def inspect(path: str) -> dict:
    """读一份 docx：页面、样式、占位符、字体依赖。"""
    d = docx.Document(path)
    sec = d.sections[0] if d.sections else None
    params = {}
    if sec:
        params = {
            "page_mm": [round(sec.page_width.mm, 1), round(sec.page_height.mm, 1)],
            "margins_mm": {
                "top": round(sec.top_margin.mm, 1),
                "bottom": round(sec.bottom_margin.mm, 1),
                "left": round(sec.left_margin.mm, 1),
                "right": round(sec.right_margin.mm, 1),
            },
        }
        params["text_width_mm"] = round(sec.page_width.mm - sec.left_margin.mm - sec.right_margin.mm, 1)
        params["text_height_mm"] = round(sec.page_height.mm - sec.top_margin.mm - sec.bottom_margin.mm, 1)

    styles = {}
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        try:
            st = d.styles[name]
        except KeyError:
            continue
        f = st.font
        pf = st.paragraph_format
        styles[name] = {
            "size_pt": round(f.size.pt, 1) if f.size else None,
            "cjk_font": _east_asian(st.element) if hasattr(st, "element") else None,
            # 注意：python-docx 的 line_spacing 在"固定值"时返回 Length（int 子类，单位 EMU），
            # 直接 isinstance(int) 会打印出 355600 这种数，必须按 .pt 转换。
            "line_spacing": _length_pt(pf.line_spacing),
            "first_line_indent_mm": round(pf.first_line_indent.mm, 1) if pf.first_line_indent else None,
        }

    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                texts.append(c.text)
    for s in d.sections:
        for part in (s.header, s.footer):
            try:
                texts += [p.text for p in part.paragraphs]
                for t in part.tables:
                    for r in t.rows:
                        for c in r.cells:
                            texts.append(c.text)
            except Exception:
                pass

    placeholders = sorted({m.group(1) for t in texts for m in PLACEHOLDER.finditer(t)})
    fonts = sorted({v["cjk_font"] for v in styles.values() if v.get("cjk_font")})
    return {
        "params": params,
        "styles": styles,
        "placeholders": placeholders,
        "fonts": fonts,
        "paragraphs": len(d.paragraphs),
        "tables": len(d.tables),
    }


def _length_pt(value):
    """Length / float / None → pt 数值；倍数行距返回原值。"""
    if value is None:
        return None
    if hasattr(value, "pt"):
        return round(value.pt, 1)
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    return None


def _east_asian(style_element) -> str | None:
    """从样式的 rPr/rFonts 里取东亚字体名（w:eastAsia）。"""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    rpr = style_element.find(f".//{ns}rPr")
    if rpr is None:
        return None
    fonts = rpr.find(f"{ns}rFonts")
    if fonts is None:
        return None
    return fonts.get(f"{ns}eastAsia")


# ─────────────────────────── 模板库 ───────────────────────────

def template_dir(name: str, store: str) -> str:
    return os.path.join(store, name)


def load_meta(name: str, store: str) -> dict:
    path = os.path.join(template_dir(name, store), "meta.json")
    if not os.path.exists(path):
        raise SystemExit(f"没有这个模板：{name}\n  已有：{', '.join(os.listdir(store)) if os.path.isdir(store) else '（空）'}")
    return json.load(open(path, encoding="utf-8"))


def cmd_template_add(args) -> int:
    store = args.store or default_store()
    src = args.file
    if not os.path.exists(src):
        raise SystemExit(f"模板不存在：{src}")
    name = args.name or os.path.splitext(os.path.basename(src))[0]
    dest_dir = template_dir(name, store)
    if os.path.exists(dest_dir) and not args.force:
        raise SystemExit(f"模板已存在：{dest_dir}（要覆盖请加 --force）")

    info = inspect(src)
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, "template.docx")
    shutil.copy2(src, dest)

    font_status = {f: font_available(f) for f in info["fonts"]}
    meta = {
        "name": name,
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source": os.path.abspath(src),
        "sha256": sha256(dest),
        **info,
        "font_status": font_status,
    }
    with open(os.path.join(dest_dir, "meta.json"), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    print(f"✓ 已存储模板：{name}")
    print(f"  路径：{dest}")
    print(f"  版心：{info['params'].get('text_width_mm')}×{info['params'].get('text_height_mm')}mm"
          f"  页边距：{info['params']['margins_mm']}")
    print(f"  占位符：{', '.join(info['placeholders']) or '（没有找到 {{…}} 占位符）'}")
    print(f"  字体：{', '.join(info['fonts']) or '（样式里没有显式东亚字体）'}")
    for f, ok in font_status.items():
        if ok is False:
            print(f"  ⚠ 本机没有字体「{f}」：Word 打开时会静默替换 —— 渲染机必须装，或换模板字体")
        elif ok is None:
            print(f"  ？ 无法在本机确认字体「{f}」是否存在（请在渲染机确认）")
    if not info["placeholders"]:
        print("  ⚠ 没找到占位符：这份模板只能当参考，render 会失败。请把可变处改成 {{字段}}")
    return 0


def cmd_template_list(args) -> int:
    store = args.store or default_store()
    if not os.path.isdir(store):
        print("（模板库为空）")
        return 0
    rows = []
    for name in sorted(os.listdir(store)):
        try:
            m = load_meta(name, store)
        except SystemExit:
            continue
        rows.append((name, m["created_at"][:10], len(m["placeholders"]), "、".join(m["fonts"])[:24]))
    if not rows:
        print("（模板库为空）")
        return 0
    print(f"{'名称':<16}{'加入日期':<12}{'占位符':<8}字体")
    for r in rows:
        print(f"{r[0]:<16}{r[1]:<12}{r[2]:<8}{r[3]}")
    return 0


def cmd_template_show(args) -> int:
    store = args.store or default_store()
    m = load_meta(args.name, store)
    print(f"模板：{m['name']}（{m['created_at'][:10]}）")
    print(f"  sha256：{m['sha256'][:16]}…")
    print(f"  页面：{m['params'].get('page_mm')}　版心：{m['params'].get('text_width_mm')}×{m['params'].get('text_height_mm')}mm")
    print(f"  页边距(mm)：{m['params']['margins_mm']}")
    print("  样式：")
    for k, v in m["styles"].items():
        indent = f"{v['first_line_indent_mm']}mm" if v["first_line_indent_mm"] is not None else "—"
        print(f"    {k:<12} 字号={v['size_pt']}  字体={v['cjk_font']}  行距={v['line_spacing']}  首行缩进={indent}")
    print(f"  占位符：{', '.join(m['placeholders']) or '（无）'}")
    if args.json:
        print(json.dumps(m, ensure_ascii=False, indent=2))
    return 0


# ─────────────────────────── 渲染 ───────────────────────────

def _set_paragraph_text(p, text: str) -> None:
    """保留第一个 run 的格式，把整段替换成 text。"""
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def _iter_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p
    for s in d.sections:
        for part in (s.header, s.footer):
            try:
                for p in part.paragraphs:
                    yield p
                for t in part.tables:
                    for r in t.rows:
                        for c in r.cells:
                            for p in c.paragraphs:
                                yield p
            except Exception:
                pass


def _insert_after(paragraph, text: str):
    """在段落之后插入一段，复制它的段落格式与首个 run 的字符格式。"""
    new_p = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, paragraph._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    run = np.add_run(text)
    if paragraph.runs:
        src = paragraph.runs[0]
        run.bold = src.bold
        run.italic = src.italic
        if src.font.size:
            run.font.size = src.font.size
        if src.font.name:
            run.font.name = src.font.name
        # 东亚字体在 rPr/rFonts 上，复制一份
        from copy import deepcopy as _dc
        src_rpr = src._r.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if src_rpr is not None:
            run._r.insert(0, _dc(src_rpr))
    return np


def cmd_render(args) -> int:
    store = args.store or default_store()
    m = load_meta(args.template, store)
    tpl_path = os.path.join(template_dir(args.template, store), "template.docx")

    if sha256(tpl_path) != m["sha256"] and not args.allow_drift:
        raise SystemExit("模板文件与登记时不一致（hash 变了）。确认是你要的版本后加 --allow-drift 重跑。")

    data = {}
    if args.data:
        with open(args.data, encoding="utf-8") as fh:
            data = json.load(fh)
    for pair in args.field or []:
        if "=" not in pair:
            raise SystemExit(f"--field 需要 k=v 形式：{pair}")
        k, v = pair.split("=", 1)
        data[k.strip()] = v
    if args.body:
        with open(args.body, encoding="utf-8") as fh:
            data.setdefault("正文", fh.read())

    unknown = sorted(set(data) - set(m["placeholders"]))
    if args.strict and unknown:
        raise SystemExit(f"数据里有模板未声明的字段：{', '.join(unknown)}（拼错了？去掉 --strict 可忽略）")

    required = [p for p in m["placeholders"] if p not in data]
    if required and not args.allow_missing:
        raise SystemExit(f"缺少必填占位符：{', '.join(required)}\n  （确实要留空请加 --allow-missing）")

    d = docx.Document(tpl_path)
    inserted = 0
    for p in list(_iter_paragraphs(d)):
        matches = list(PLACEHOLDER.finditer(p.text))
        if not matches:
            continue
        if len(matches) == 1 and matches[0].group(1) in data:
            key = matches[0].group(1)
            value = str(data[key])
            whole = p.text.strip() == matches[0].group(0)
            if whole and "\n" in value:
                # 整段就是占位符、且内容是多段：拆成多段，保留样式。
                # 空行丢掉：公文靠"首行缩进 2 字"分段，不用空行分隔（空行会让版式散掉）。
                lines = [ln for ln in value.split("\n") if ln.strip()]
                if not lines:
                    _set_paragraph_text(p, "")
                    continue
                _set_paragraph_text(p, lines[0])
                anchor = p
                for line in lines[1:]:
                    anchor = _insert_after(anchor, line)
                    inserted += 1
                continue
            _set_paragraph_text(p, p.text.replace(matches[0].group(0), value))
        else:
            text = p.text
            for mt in matches:
                key = mt.group(1)
                if key in data:
                    text = text.replace(mt.group(0), str(data[key]))
            _set_paragraph_text(p, text)

    out = args.out
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)

    # 生成前的最后一道：产物里不得残留 {{占位符}}。
    # 交付一份带 {{姓名}} 的公文比报错难看得多，而它"看起来是成功的"。
    left = sorted({m.group(1) for p in _iter_paragraphs(d) for m in PLACEHOLDER.finditer(p.text)})
    if left and args.strict:
        # strict 下不写文件；留一份带占位符的成文比报错更难发现
        raise SystemExit(f"仍残留占位符：{', '.join(left)}（strict 模式不生成文件）")
    d.save(out)
    print(f"✓ 已生成：{out}（补了 {inserted} 段）")
    if left:
        print(f"  ⚠ 仍有 {len(left)} 个占位符未替换：{', '.join(left)}（交付前必须处理）")
    print(f"  模板：{args.template}（sha256 {m['sha256'][:12]}…）")

    # 生成即自检：与模板登记参数比对
    problems = compare(out, m)
    if problems:
        print("  ⚠ 产物与模板参数不一致：")
        for x in problems:
            print(f"    - {x}")
        return 1
    print("  ✓ 产物与模板参数一致")
    return 0


def compare(docx_path: str, meta: dict) -> list[str]:
    info = inspect(docx_path)
    exp = meta["params"]
    got = info["params"]
    problems = []
    for side in ("top", "bottom", "left", "right"):
        a, b = exp["margins_mm"][side], got["margins_mm"][side]
        if abs(a - b) > 0.6:
            problems.append(f"页边距 {side}: 模板 {a}mm → 产物 {b}mm")
    if exp.get("page_mm") != got.get("page_mm"):
        problems.append(f"页面尺寸: 模板 {exp.get('page_mm')} → 产物 {got.get('page_mm')}")
    for name in ("Normal", "Heading 1", "Heading 2"):
        e = meta["styles"].get(name) or {}
        g = info["styles"].get(name) or {}
        if e.get("size_pt") and g.get("size_pt") and abs(e["size_pt"] - g["size_pt"]) > 0.2:
            problems.append(f"{name} 字号: 模板 {e['size_pt']}pt → 产物 {g['size_pt']}pt")
        # 行距与首行缩进是公文最容易被"顺手改坏"的两项，必须一起量
        if e.get("line_spacing") and g.get("line_spacing") and abs(e["line_spacing"] - g["line_spacing"]) > 0.5:
            problems.append(f"{name} 行距: 模板 {e['line_spacing']}pt → 产物 {g['line_spacing']}pt")
        if e.get("first_line_indent_mm") is not None and g.get("first_line_indent_mm") is not None \
                and abs(e["first_line_indent_mm"] - g["first_line_indent_mm"]) > 0.6:
            problems.append(f"{name} 首行缩进: 模板 {e['first_line_indent_mm']}mm → 产物 {g['first_line_indent_mm']}mm")
    return problems


def cmd_check(args) -> int:
    store = args.store or default_store()
    m = load_meta(args.template, store)
    problems = compare(args.file, m)
    if not problems:
        print(f"✓ {args.file} 与模板「{args.template}」参数一致")
        return 0
    print(f"✗ {args.file} 与模板「{args.template}」不一致：")
    for p in problems:
        print(f"  - {p}")
    return 1


# ─────────────────────────── CLI ───────────────────────────

def main(argv=None) -> int:
    from . import __version__
    ap = argparse.ArgumentParser(description="公文模板引擎（模板固化，不生成红头）")
    ap.add_argument("--version", action="version", version=f"md-to-office {__version__}（gongwen）")
    ap.add_argument("--store", help="模板库目录（默认 $GONGWEN_HOME 或 ~/.md-to-office/templates）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    t = sub.add_parser("template", help="模板管理")
    tsub = t.add_subparsers(dest="sub", required=True)
    a = tsub.add_parser("add", help="存储并体检一份模板")
    a.add_argument("file")
    a.add_argument("--name")
    a.add_argument("--force", action="store_true")
    a.set_defaults(func=cmd_template_add)
    s = tsub.add_parser("show", help="查看模板参数")
    s.add_argument("name")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_template_show)
    l = tsub.add_parser("list", help="列出模板")
    l.set_defaults(func=cmd_template_list)

    r = sub.add_parser("render", help="按模板填充生成")
    r.add_argument("--template", required=True)
    r.add_argument("--data", help="JSON：{占位符: 值}")
    r.add_argument("--field", action="append", help="k=v，可重复")
    r.add_argument("--body", help="正文文件（Markdown 纯文本，填入 {{正文}}）")
    r.add_argument("--out", required=True)
    r.add_argument("--allow-missing", action="store_true")
    r.add_argument("--allow-drift", action="store_true")
    r.add_argument("--strict", action="store_true",
                   help="未知字段或残留占位符一律报错（不生成文件）")
    r.set_defaults(func=cmd_render)

    c = sub.add_parser("check", help="校验产物是否还符合模板")
    c.add_argument("file")
    c.add_argument("--template", required=True)
    c.set_defaults(func=cmd_check)

    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
