#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown → DOCX 引擎（可配置、可复用）

用法（库）：
    from md2docx import Config, convert
    convert("in.md", "out.docx", Config(body_ea="宋体", heading_ea="微软雅黑"))

用法（CLI）：
    python md2docx.py --input in.md --output out.docx [--font-body 宋体] ...

支持语法：标题 1–4 级、表格（表头跨页重复 + 行不跨页断开）、无序/有序列表、
引用块、代码块、独占一行的图片、粗体、行内代码、链接、分隔线。
"""
from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ─────────────────────── 配置 ───────────────────────


@dataclass
class Config:
    """排版配置。默认值面向中文商务文档（A4 + 宋体正文 + 微软雅黑标题）。"""

    # 字体
    body_ascii: str = "Times New Roman"
    body_ea: str = "宋体"
    heading_ascii: str = "Microsoft YaHei"
    heading_ea: str = "微软雅黑"
    code_font: str = "Consolas"

    # 字号（pt）
    body_size: float = 10.5
    h1_size: float = 15.0
    h2_size: float = 13.0
    h3_size: float = 11.5
    h4_size: float = 10.5
    table_size: float = 9.0
    code_size: float = 8.5
    footer_size: float = 9.0

    # 页面（cm）
    page_width: float = 21.0          # A4 宽；Letter 用 21.59
    page_height: float = 29.7         # A4 高；Letter 用 27.94
    margin_top: float = 2.2
    margin_bottom: float = 2.2
    margin_left: float = 2.0
    margin_right: float = 2.0

    # 颜色
    h1_color: tuple = (0x1F, 0x38, 0x64)
    h2_color: tuple = (0x1F, 0x38, 0x64)
    h3_color: tuple = (0x2E, 0x54, 0x96)
    h4_color: tuple = (0x2E, 0x54, 0x96)
    table_header_fill: str = "DEEAF6"
    link_color: tuple = (0x1F, 0x4E, 0x79)
    code_color: tuple = (0xC0, 0x39, 0x2B)
    footer_color: tuple = (0x80, 0x80, 0x80)

    # 段落
    line_spacing: float = 1.15
    space_after: float = 4.0

    # 开关
    page_number_footer: bool = True
    repeat_table_header: bool = True
    keep_table_rows_together: bool = True

    @property
    def usable_width_cm(self) -> float:
        return self.page_width - self.margin_left - self.margin_right

    def apply_overrides(self, **kw) -> "Config":
        for k, v in kw.items():
            if v is not None and hasattr(self, k):
                setattr(self, k, v)
        return self


# ─────────────────────── 低层工具 ───────────────────────


def _rgb(t):
    return RGBColor(*t) if isinstance(t, (tuple, list)) else t


def set_run(run, cfg: Config, ascii_font=None, ea_font=None, size=None,
            bold=None, italic=None, color=None):
    af = ascii_font or cfg.body_ascii
    ef = ea_font or cfg.body_ea
    run.font.name = af
    run.font.size = Pt(size if size is not None else cfg.body_size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = _rgb(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), af)
    rFonts.set(qn("w:hAnsi"), af)
    rFonts.set(qn("w:eastAsia"), ef)
    rFonts.set(qn("w:cs"), af)


def style_font(style, ascii_font, ea_font, size, bold=False, color=None):
    style.font.name = ascii_font
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = _rgb(color)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), ea_font)
    rFonts.set(qn("w:cs"), ascii_font)


def shade_cell(tc, fill):
    pr = tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    pbdr.append(bottom)
    pPr.append(pbdr)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_page_number_footer(section, cfg: Config):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    set_run(r, cfg, size=cfg.footer_size, color=cfg.footer_color)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


# ─────────────────────── 行内解析 ───────────────────────

TOKEN = re.compile(r"(!\[[^\]]*\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, cfg: Config, size=None, base_bold=False,
               base_italic=False, base_color=None, mono_all=False):
    size = size if size is not None else cfg.body_size
    text = text.replace("\u200b", "")
    if mono_all:
        r = paragraph.add_run(text)
        set_run(r, cfg, ascii_font=cfg.code_font, ea_font=cfg.code_font,
                size=size, bold=base_bold, italic=base_italic, color=base_color)
        return
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("![") and "](" in part:
            continue                      # 图片在块级处理
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            set_run(r, cfg, size=size, bold=True, italic=base_italic, color=base_color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            set_run(r, cfg, ascii_font=cfg.code_font, ea_font=cfg.code_font,
                    size=max(size - 0.5, 6), bold=base_bold, italic=base_italic,
                    color=cfg.code_color)
        elif part.startswith("[") and "](" in part:
            label = part[1:part.index("](")]
            r = paragraph.add_run(label)
            set_run(r, cfg, size=size, bold=base_bold, italic=base_italic,
                    color=cfg.link_color)
            r.font.underline = True
        else:
            r = paragraph.add_run(part)
            set_run(r, cfg, size=size, bold=base_bold, italic=base_italic,
                    color=base_color)


# ─────────────────────── 转换主体 ───────────────────────

_H_SIZE = ("h1_size", "h2_size", "h3_size", "h4_size")


def convert(md_path: str, docx_path: str, cfg: Config | None = None,
            image_base: str | None = None, verbose: bool = False) -> str:
    """把单个 Markdown 转成 DOCX，返回输出路径。"""
    cfg = cfg or Config()
    raw = open(md_path, encoding="utf-8").read().replace("\r\n", "\n")
    lines = raw.split("\n")
    img_root = image_base or os.path.dirname(os.path.abspath(md_path))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(cfg.page_width), Cm(cfg.page_height)
    sec.top_margin, sec.bottom_margin = Cm(cfg.margin_top), Cm(cfg.margin_bottom)
    sec.left_margin, sec.right_margin = Cm(cfg.margin_left), Cm(cfg.margin_right)
    if cfg.page_number_footer:
        add_page_number_footer(sec, cfg)

    st = doc.styles
    style_font(st["Normal"], cfg.body_ascii, cfg.body_ea, cfg.body_size)
    for lvl, key, col in ((1, "h1_size", cfg.h1_color), (2, "h2_size", cfg.h2_color),
                          (3, "h3_size", cfg.h3_color), (4, "h4_size", cfg.h4_color)):
        style_font(st[f"Heading {lvl}"], cfg.heading_ascii, cfg.heading_ea,
                   getattr(cfg, key), True, col)
    st["Normal"].paragraph_format.space_after = Pt(cfg.space_after)
    st["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    st["Normal"].paragraph_format.line_spacing = cfg.line_spacing

    usable = Cm(cfg.usable_width_cm)
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]
        s = line.strip()

        # 图片（独占一行）
        mi = re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", s)
        if mi:
            ip = mi.group(1)
            if not os.path.isabs(ip):
                ip = os.path.normpath(os.path.join(img_root, ip))
            if os.path.exists(ip):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                p.add_run().add_picture(ip, width=usable)
                if verbose:
                    print("      [图]", os.path.basename(ip))
            else:
                print("      [图缺失]", ip, file=sys.stderr)
            i += 1
            continue

        # 代码块
        if s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            for bl in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                add_inline(p, bl if bl else " ", cfg, size=cfg.code_size, mono_all=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 表格
        if s.startswith("|") and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            ncol = len(header)
            tbl = doc.add_table(rows=1, cols=ncol)
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            w = int(usable / ncol)
            hdr = tbl.rows[0]
            if cfg.repeat_table_header:
                repeat_header(hdr)
            if cfg.keep_table_rows_together:
                no_split(hdr)
            for c, txt in enumerate(header):
                cell = hdr.cells[c]
                cell.width = w
                shade_cell(cell._tc, cfg.table_header_fill)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                add_inline(p, txt, cfg, size=cfg.table_size, base_bold=True)
            for r in rows:
                newrow = tbl.add_row()
                if cfg.keep_table_rows_together:
                    no_split(newrow)
                cells = newrow.cells
                for c in range(ncol):
                    cell = cells[c]
                    cell.width = w
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    add_inline(p, r[c] if c < len(r) else "", cfg, size=cfg.table_size)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 分隔线
        if s in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            add_bottom_border(p)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = min(len(m.group(1)), 4)
            p = doc.add_heading(level=lvl)
            p.paragraph_format.space_before = Pt(10 if lvl <= 2 else 6)
            p.paragraph_format.space_after = Pt(5)
            col = (cfg.h1_color, cfg.h2_color, cfg.h3_color, cfg.h4_color)[lvl - 1]
            add_inline(p, m.group(2), cfg, size=getattr(cfg, _H_SIZE[lvl - 1]),
                       base_bold=True, base_color=col)
            i += 1
            continue

        # 引用块
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            for bl in buf:
                if not bl:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, bl, cfg, size=cfg.body_size, base_color=(0x44, 0x44, 0x44))
            continue

        # 列表
        mb = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        mn = re.match(r"^(\s*)(\d+)[.、]\s+(.*)$", line)
        if mb or mn:
            indent = len((mb or mn).group(1)) // 2
            txt = "• " + mb.group(2) if mb else f"{mn.group(2)}. {mn.group(3)}"
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55 + 0.45 * indent)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt, cfg)
            i += 1
            continue

        if not s:
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(cfg.space_after)
        add_inline(p, s, cfg)
        i += 1

    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(docx_path)
    return docx_path


# ─────────────────────── CLI ───────────────────────

_FONT_KEYS = {"font-body": "body_ea", "font-heading": "heading_ea",
              "font-code": "code_font", "font-body-ascii": "body_ascii"}
_SIZE_KEYS = {"size-body": "body_size", "size-table": "table_size"}


def main(argv=None):
    import argparse

    ap = argparse.ArgumentParser(description="Markdown → DOCX（可配置）")
    ap.add_argument("--input", required=True, help="输入 .md 文件")
    ap.add_argument("--output", required=True, help="输出 .docx 文件")
    ap.add_argument("--font-body", help="正文中文字体（默认 宋体）")
    ap.add_argument("--font-heading", help="标题中文字体（默认 微软雅黑）")
    ap.add_argument("--font-code", help="代码字体（默认 Consolas）")
    ap.add_argument("--size-body", type=float, help="正文字号 pt（默认 10.5）")
    ap.add_argument("--size-table", type=float, help="表格字号 pt（默认 9）")
    ap.add_argument("--page", choices=["a4", "letter"], default="a4", help="纸张")
    ap.add_argument("--no-footer", action="store_true", help="不显示页码页脚")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(argv)

    cfg = Config()
    cfg.apply_overrides(**{_FONT_KEYS[k]: v for k, v in
                           (("font-body", a.font_body), ("font-heading", a.font_heading),
                            ("font-code", a.font_code)) if v},
                        **{_SIZE_KEYS[k]: v for k, v in
                           (("size-body", a.size_body), ("size-table", a.size_table)) if v})
    if a.page == "letter":
        cfg.page_width, cfg.page_height = 21.59, 27.94
    if a.no_footer:
        cfg.page_number_footer = False

    convert(a.input, a.output, cfg, verbose=a.verbose)
    print("  ✓", a.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
