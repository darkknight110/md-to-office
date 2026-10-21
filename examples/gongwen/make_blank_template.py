#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成一份**空白公文模板**（无红头、无单位信息），用来试模板引擎。

为什么用脚本而不是提交一个 .docx：
  1. 参数可读 —— 版心/字号/行距都写在代码里，评审能一眼核对；
  2. 可复跑 —— CI 里现场生成，不依赖仓库里的二进制；
  3. 不冒充真模板 —— 它只是"符合 GB/T 9704 版式的空白页"，红头/版记要你自己单位的。

用法：
    python3 examples/gongwen/make_blank_template.py --out /tmp/空白公文.docx
"""
from __future__ import annotations

import argparse

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

# GB/T 9704-2012：A4，天头 37 / 下白边 35 / 订口 28 / 右白边 26，版心 156×225
MARGINS = {"top": 37, "bottom": 35, "left": 28, "right": 26}
BODY_PT = 16          # 3 号字
LINE_PT = 28          # 行距 28 磅（固定值）
INDENT_PT = 32        # 首行缩进 2 字符 = 2 × 16pt


def set_cjk(style, font: str) -> None:
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)


def build(path: str) -> str:
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(MARGINS["top"]), Mm(MARGINS["bottom"])
    sec.left_margin, sec.right_margin = Mm(MARGINS["left"]), Mm(MARGINS["right"])

    normal = d.styles["Normal"]
    normal.font.size = Pt(BODY_PT)
    normal.font.name = "仿宋"
    set_cjk(normal, "仿宋")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_PT)
    pf.first_line_indent = Pt(INDENT_PT)

    h1 = d.styles["Heading 1"]
    h1.font.size = Pt(22)          # 2 号
    set_cjk(h1, "方正小标宋简体")
    h1.paragraph_format.first_line_indent = Pt(0)

    d.add_paragraph("{{标题}}", style="Heading 1")
    d.add_paragraph("{{主送机关}}")
    d.add_paragraph("{{正文}}")
    d.add_paragraph("{{附件说明}}")
    d.add_paragraph("{{成文日期}}")
    # 印章位只保留位置，绝不生成图：印章是物理/电子用印，不能由排版工具产生
    d.add_paragraph("〔印章位·不自动生成〕")
    d.add_paragraph("")
    d.add_paragraph("{{版记}}")
    d.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    build(args.out)
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
